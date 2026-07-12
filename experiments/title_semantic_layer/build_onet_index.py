from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from .common import normalize_title, repo_root, write_json, write_jsonl


try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


def _read_xlsx_rows(path: Path) -> list[dict[str, str]]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read O*NET .xlsx files. Run with: uv run --with openpyxl ...")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h or "").strip() for h in next(rows_iter)]
    out: list[dict[str, str]] = []
    for row in rows_iter:
        rec = {headers[i]: str(row[i] or "").strip() for i in range(min(len(headers), len(row)))}
        out.append(rec)
    wb.close()
    return out


def build_onet_index(onet_dir: Path, out_jsonl: Path, out_meta: Path) -> dict:
    occ_path = onet_dir / "Occupation Data.xlsx"
    alt_path = onet_dir / "Alternate Titles.xlsx"
    rel_path = onet_dir / "Related Occupations.xlsx"

    occ_rows = _read_xlsx_rows(occ_path)
    alt_rows = _read_xlsx_rows(alt_path)
    rel_rows = _read_xlsx_rows(rel_path)

    alt_by_code: dict[str, list[str]] = defaultdict(list)
    for row in alt_rows:
        code = row.get("O*NET-SOC Code", "")
        alt_title = row.get("Alternate Title", "")
        if code and alt_title:
            alt_by_code[code].append(alt_title)

    rel_by_code: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rel_rows:
        code = row.get("O*NET-SOC Code", "")
        rel_title = row.get("Related Title", "")
        rel_code = row.get("Related O*NET-SOC Code", "")
        tier = row.get("Relatedness Tier", "")
        if code and rel_title:
            rel_by_code[code].append(
                {
                    "related_title": rel_title,
                    "related_code": rel_code,
                    "relatedness_tier": tier,
                }
            )

    out_rows = []
    for row in occ_rows:
        code = row.get("O*NET-SOC Code", "")
        title = row.get("Title", "")
        desc = row.get("Description", "")
        if not code or not title:
            continue
        out_rows.append(
            {
                "source": "onet",
                "soc_code": code,
                "title": title,
                "normalized_title": normalize_title(title),
                "description": desc,
                "alternate_titles": alt_by_code.get(code, []),
                "related_occupations": rel_by_code.get(code, [])[:8],
            }
        )

    write_jsonl(out_jsonl, out_rows)
    meta = {
        "occupation_rows": len(occ_rows),
        "alternate_rows": len(alt_rows),
        "related_rows": len(rel_rows),
        "index_rows": len(out_rows),
        "files": {
            "Occupation Data": str(occ_path),
            "Alternate Titles": str(alt_path),
            "Related Occupations": str(rel_path),
        },
    }
    write_json(out_meta, meta)
    return meta


if __name__ == "__main__":
    root = repo_root()
    onet = root / "ONETfiles/db_30_2_excel"
    out_jsonl = root / "experiments/title_semantic_layer/reports/onet_index.jsonl"
    out_meta = root / "experiments/title_semantic_layer/reports/onet_index_meta.json"
    m = build_onet_index(onet, out_jsonl, out_meta)
    print(f"onet_index -> {out_jsonl} ({m['index_rows']} rows)")
