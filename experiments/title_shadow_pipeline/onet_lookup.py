from __future__ import annotations

import argparse
import json
from pathlib import Path

from common import normalize_title

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


def _iter_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(min_row=1, values_only=True)
    header = [str(x or "").strip() for x in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    for row in rows:
        yield idx, row
    wb.close()


def run(onet_dir: Path, outdir: Path) -> dict[str, object]:
    if load_workbook is None:
        raise RuntimeError("openpyxl not available; run with: uv run --with openpyxl ...")

    base = onet_dir / "db_30_2_excel"
    occ_path = base / "Occupation Data.xlsx"
    alt_path = base / "Alternate Titles.xlsx"
    rel_path = base / "Related Occupations.xlsx"

    entries: list[dict[str, str]] = []

    for idx, row in _iter_rows(occ_path):
        title = str(row[idx.get("Title", -1)] or "").strip()
        if not title:
            continue
        code = str(row[idx.get("O*NET-SOC Code", -1)] or "").strip()
        n = normalize_title(title)
        if not n:
            continue
        entries.append(
            {
                "source": "onet",
                "normalized_title": n,
                "raw_title": title,
                "canonical_title": title,
                "onet_code": code,
                "relation": "occupation_title",
            }
        )

    for idx, row in _iter_rows(alt_path):
        title = str(row[idx.get("Alternate Title", -1)] or "").strip()
        canonical = str(row[idx.get("Title", -1)] or "").strip()
        code = str(row[idx.get("O*NET-SOC Code", -1)] or "").strip()
        if not title:
            continue
        n = normalize_title(title)
        if not n:
            continue
        entries.append(
            {
                "source": "onet",
                "normalized_title": n,
                "raw_title": title,
                "canonical_title": canonical or title,
                "onet_code": code,
                "relation": "alternate_title",
            }
        )

    for idx, row in _iter_rows(rel_path):
        title = str(row[idx.get("Related Title", -1)] or "").strip()
        canonical = str(row[idx.get("Title", -1)] or "").strip()
        code = str(row[idx.get("Related O*NET-SOC Code", -1)] or "").strip()
        if not title:
            continue
        n = normalize_title(title)
        if not n:
            continue
        entries.append(
            {
                "source": "onet",
                "normalized_title": n,
                "raw_title": title,
                "canonical_title": canonical or title,
                "onet_code": code,
                "relation": "related_occupation",
            }
        )

    # Deduplicate on (normalized_title, relation, canonical_title)
    uniq = {}
    for e in entries:
        key = (e["normalized_title"], e["relation"], e["canonical_title"])
        uniq[key] = e
    rows = list(uniq.values())

    outdir.mkdir(parents=True, exist_ok=True)
    out_path = outdir / "onet_lookup.jsonl"
    with out_path.open("w", encoding="utf-8") as f:
        for e in rows:
            f.write(json.dumps(e, ensure_ascii=False) + "\n")

    meta = {
        "lookup_rows": len(rows),
        "files": {
            "Occupation Data": str(occ_path),
            "Alternate Titles": str(alt_path),
            "Related Occupations": str(rel_path),
        },
    }
    (outdir / "onet_lookup_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return meta


def main() -> None:
    ap = argparse.ArgumentParser(description="Build O*NET lookup index")
    ap.add_argument("--onet-dir", required=True)
    ap.add_argument("--outdir", required=True)
    args = ap.parse_args()
    run(Path(args.onet_dir), Path(args.outdir))


if __name__ == "__main__":
    main()
